"""
Тесты document_extractor: DOCX, XLSX и ZIP создаются на лету (python-docx, openpyxl, zipfile)
во временной директории или в памяти. Сеть не используется.

Запуск из корня проекта:
    python -m unittest tests.test_document_extractor -v
"""

import datetime
import io
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import Workbook

from src.parser import document_extractor as extractor

ARMENIAN = "Գնման առարկա՝ ապրանքների ուղեփոխադրում"
RUSSIAN = "Предмет закупки: международные грузоперевозки"

BROKEN_DOCX_ERRORS = (PackageNotFoundError, zipfile.BadZipFile)


def make_docx_bytes(paragraphs=(), tables=()):
    """tables — список таблиц, таблица — список строк, строка — список ячеек."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(path, members, compression=zipfile.ZIP_STORED):
    """members — список (имя, bytes); имя, оканчивающееся на '/', — директория."""
    with zipfile.ZipFile(path, "w", compression) as zf:
        for name, data in members:
            zf.writestr(name, data)


def make_zip_bytes(members, compression=zipfile.ZIP_STORED):
    buffer = io.BytesIO()
    make_zip(buffer, members, compression)
    return buffer.getvalue()


def make_xlsx_bytes(sheets):
    """sheets — {название листа: список строк}, строка — список значений ячеек."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def replace_zip_member(data, member_name, transform):
    """Копия ZIP (xlsx/docx), в которой содержимое одного member изменено функцией transform."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == member_name:
                content = transform(content)
            target.writestr(info, content)
    return out.getvalue()


class ExtractDocxTests(unittest.TestCase):
    def test_single_paragraph(self):
        result = extractor.extract_docx(make_docx_bytes(["Hello tender"]))
        self.assertEqual(result["file_type"], "docx")
        self.assertEqual(result["text"], "Hello tender")
        self.assertEqual(result["paragraph_count"], 1)
        self.assertEqual(result["table_count"], 0)

    def test_multiple_paragraphs(self):
        result = extractor.extract_docx(make_docx_bytes(["first", "second", "third"]))
        self.assertEqual(result["text"], "first\nsecond\nthird")
        self.assertEqual(result["paragraph_count"], 3)

    def test_empty_paragraphs_do_not_create_garbage(self):
        result = extractor.extract_docx(make_docx_bytes(["first", "", "   ", "second"]))
        self.assertEqual(result["text"], "first\nsecond")
        self.assertEqual(result["paragraph_count"], 2)

    def test_armenian_unicode_preserved(self):
        result = extractor.extract_docx(make_docx_bytes([ARMENIAN]))
        self.assertEqual(result["text"], ARMENIAN)

    def test_russian_unicode_preserved(self):
        result = extractor.extract_docx(make_docx_bytes([RUSSIAN]))
        self.assertEqual(result["text"], RUSSIAN)

    def test_inner_whitespace_not_normalized(self):
        result = extractor.extract_docx(make_docx_bytes(["  a  b   c  "]))
        self.assertEqual(result["text"], "a  b   c")

    def test_table_extracted(self):
        result = extractor.extract_docx(make_docx_bytes(tables=[[["A", "B", "C"]]]))
        self.assertEqual(result["text"], "A\tB\tC")
        self.assertEqual(result["table_count"], 1)
        self.assertEqual(result["paragraph_count"], 0)

    def test_table_multiple_rows(self):
        rows = [["Name", "Qty"], ["Bolt", "10"], [ARMENIAN, RUSSIAN]]
        result = extractor.extract_docx(make_docx_bytes(tables=[rows]))
        self.assertEqual(result["text"], f"Name\tQty\nBolt\t10\n{ARMENIAN}\t{RUSSIAN}")

    def test_paragraph_and_table_both_present(self):
        result = extractor.extract_docx(
            make_docx_bytes(["Intro paragraph"], tables=[[["k", "v"]]])
        )
        self.assertIn("Intro paragraph", result["text"])
        self.assertIn("k\tv", result["text"])

    def test_table_with_empty_row_skipped(self):
        result = extractor.extract_docx(make_docx_bytes(tables=[[["a", "b"], ["", ""]]]))
        self.assertEqual(result["text"], "a\tb")

    def test_char_count_matches_text_length(self):
        result = extractor.extract_docx(
            make_docx_bytes([ARMENIAN, RUSSIAN], tables=[[["x", "y"]]])
        )
        self.assertEqual(result["char_count"], len(result["text"]))
        self.assertGreater(result["char_count"], 0)

    def test_bytes_source(self):
        result = extractor.extract_docx(make_docx_bytes(["from bytes"]))
        self.assertEqual(result["text"], "from bytes")

    def test_file_like_source(self):
        stream = io.BytesIO(make_docx_bytes(["from stream"]))
        result = extractor.extract_docx(stream)
        self.assertEqual(result["text"], "from stream")

    def test_path_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "doc.docx"
            path.write_bytes(make_docx_bytes(["from path"]))
            self.assertEqual(extractor.extract_docx(path)["text"], "from path")
            self.assertEqual(extractor.extract_docx(str(path))["text"], "from path")

    def test_empty_docx(self):
        result = extractor.extract_docx(make_docx_bytes())
        self.assertEqual(result["text"], "")
        self.assertEqual(result["char_count"], 0)
        self.assertEqual(result["paragraph_count"], 0)
        self.assertEqual(result["table_count"], 0)

    def test_broken_docx_raises(self):
        with self.assertRaises(BROKEN_DOCX_ERRORS):
            extractor.extract_docx(b"this is not a docx file")


class CleanTextPieceTests(unittest.TestCase):
    def test_strip_and_newlines(self):
        self.assertEqual(extractor.clean_text_piece("  a\r\nb\rc  "), "a\nb\nc")

    def test_inner_spaces_and_unicode_untouched(self):
        self.assertEqual(extractor.clean_text_piece(f" {ARMENIAN}  x "), f"{ARMENIAN}  x")


class ExtractDocxFromZipTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        self.zip_path = self.tmp / "tender.zip"

    def test_two_docx(self):
        make_zip(self.zip_path, [
            ("hraver.docx", make_docx_bytes([ARMENIAN])),
            ("hraver_ru.docx", make_docx_bytes([RUSSIAN], tables=[[["a", "b"]]])),
        ])
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["file_type"], "zip")
        self.assertEqual(result["member_count"], 2)
        self.assertEqual(result["docx_count"], 2)
        self.assertEqual(result["skipped_members"], [])
        self.assertEqual(result["failures"], [])

    def test_texts_of_both_docx_extracted(self):
        make_zip(self.zip_path, [
            ("hraver.docx", make_docx_bytes([ARMENIAN])),
            ("hraver_ru.docx", make_docx_bytes([RUSSIAN])),
        ])
        docs = extractor.extract_docx_from_zip(self.zip_path)["documents"]
        by_name = {d["member_name"]: d for d in docs}
        self.assertEqual(by_name["hraver.docx"]["text"], ARMENIAN)
        self.assertEqual(by_name["hraver_ru.docx"]["text"], RUSSIAN)
        self.assertEqual(by_name["hraver.docx"]["file_type"], "docx")
        self.assertEqual(by_name["hraver.docx"]["char_count"], len(ARMENIAN))

    def test_unsupported_member_skipped(self):
        make_zip(self.zip_path, [
            ("a.docx", make_docx_bytes(["a"])),
            ("b.pdf", b"%PDF-1.4"),
            ("c.txt", b"text"),
            ("d.doc", b"\xd0\xcf"),
        ])
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["docx_count"], 1)
        self.assertEqual(result["skipped_members"], ["b.pdf", "c.txt", "d.doc"])

    def test_directory_member_ignored(self):
        make_zip(self.zip_path, [
            ("folder/", b""),
            ("folder/a.docx", make_docx_bytes(["a"])),
        ])
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["member_count"], 1)
        self.assertEqual(result["docx_count"], 1)
        self.assertEqual(result["skipped_members"], [])

    def test_uppercase_extension_supported(self):
        make_zip(self.zip_path, [("HRAVER.DOCX", make_docx_bytes(["upper"]))])
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["docx_count"], 1)
        self.assertEqual(result["documents"][0]["text"], "upper")

    def test_too_many_members(self):
        make_zip(self.zip_path, [(f"{i}.txt", b"x") for i in range(3)])
        with self.assertRaises(ValueError):
            extractor.extract_docx_from_zip(self.zip_path, max_members=2)

    def test_default_max_members_is_200(self):
        self.assertEqual(extractor.DEFAULT_MAX_MEMBERS, 200)

    def test_59_safe_members_not_rejected_by_count(self):
        # Реальный eAuction ZIP содержал 59 member (2.8 MB): 2 docx + прочие файлы.
        members = [("hraver.docx", make_docx_bytes([ARMENIAN]))]
        members += [(f"file_{i}.pdf", b"%PDF-1.4") for i in range(58)]
        make_zip(self.zip_path, members)
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["member_count"], 59)
        self.assertEqual(result["docx_count"], 1)
        self.assertEqual(len(result["skipped_members"]), 58)
        self.assertEqual(result["failures"], [])

    def test_200_members_allowed_by_default(self):
        make_zip(self.zip_path, [(f"{i}.txt", b"x") for i in range(200)])
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["member_count"], 200)
        self.assertEqual(len(result["skipped_members"]), 200)

    def test_201_members_rejected_by_default(self):
        make_zip(self.zip_path, [(f"{i}.txt", b"x") for i in range(201)])
        with self.assertRaises(ValueError) as ctx:
            extractor.extract_docx_from_zip(self.zip_path)
        self.assertIn("too many members", str(ctx.exception))

    def test_directory_entries_not_counted_toward_member_limit(self):
        members = [(f"d{i}/", b"") for i in range(50)]
        members += [(f"{i}.txt", b"x") for i in range(200)]
        make_zip(self.zip_path, members)
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["member_count"], 200)

    def test_member_too_large(self):
        make_zip(self.zip_path, [("big.bin", b"x" * 100)])
        with self.assertRaises(ValueError):
            extractor.extract_docx_from_zip(self.zip_path, max_member_bytes=50)

    def test_total_uncompressed_too_large(self):
        make_zip(self.zip_path, [("a.bin", b"x" * 60), ("b.bin", b"x" * 60)])
        with self.assertRaises(ValueError):
            extractor.extract_docx_from_zip(
                self.zip_path, max_member_bytes=100, max_total_uncompressed_bytes=100
            )

    def test_broken_docx_goes_to_failures_others_extracted(self):
        make_zip(self.zip_path, [
            ("broken.docx", b"not a real docx"),
            ("good.docx", make_docx_bytes([RUSSIAN])),
        ])
        result = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(result["docx_count"], 1)
        self.assertEqual(result["documents"][0]["member_name"], "good.docx")
        self.assertEqual(result["documents"][0]["text"], RUSSIAN)
        self.assertEqual(len(result["failures"]), 1)
        failure = result["failures"][0]
        self.assertEqual(failure["member_name"], "broken.docx")
        self.assertTrue(failure["error_type"])
        self.assertIn("error_message", failure)

    def test_zip_writes_nothing_to_filesystem(self):
        work = self.tmp / "work"
        work.mkdir()
        zip_path = work / "tender.zip"
        make_zip(zip_path, [
            ("../evil.docx", make_docx_bytes(["evil"])),
            ("/abs_evil.docx", make_docx_bytes(["abs"])),
            ("ok.docx", make_docx_bytes(["ok"])),
        ])
        before_tree = sorted(str(p) for p in self.tmp.rglob("*"))
        cwd_before = sorted(os.listdir("."))

        result = extractor.extract_docx_from_zip(zip_path)

        self.assertEqual(result["docx_count"], 3)
        self.assertEqual(sorted(str(p) for p in self.tmp.rglob("*")), before_tree)
        self.assertFalse((self.tmp / "evil.docx").exists())
        self.assertEqual(sorted(os.listdir(".")), cwd_before)
        # member_name сохранён как метаданные без изменений
        names = {d["member_name"] for d in result["documents"]}
        self.assertIn("../evil.docx", names)


class ExtractXlsxTests(unittest.TestCase):
    def test_single_sheet(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({"Specs": [["Item", "Qty"], ["Bolt", "10"]]}))
        self.assertEqual(result["file_type"], "xlsx")
        self.assertEqual(result["text"], "[Sheet: Specs]\nItem\tQty\nBolt\t10")
        self.assertEqual(result["sheet_count"], 1)
        self.assertEqual(result["row_count"], 2)
        self.assertEqual(result["cell_count"], 4)

    def test_multiple_sheets(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({
            "Lot 1": [["a", "b"]],
            "Lot 2": [["c"], ["d", "e"]],
        }))
        self.assertEqual(result["text"], "[Sheet: Lot 1]\na\tb\n\n[Sheet: Lot 2]\nc\nd\te")
        self.assertEqual(result["sheet_count"], 2)
        self.assertEqual(result["row_count"], 3)
        self.assertEqual(result["cell_count"], 5)

    def test_empty_sheet_is_skipped(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({
            "Empty": [],
            "Blank": [[None, "", None]],
            "Data": [["x"]],
        }))
        self.assertEqual(result["text"], "[Sheet: Data]\nx")
        self.assertEqual(result["sheet_count"], 1)

    def test_empty_workbook(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({"Sheet": []}))
        self.assertEqual(result["text"], "")
        self.assertEqual(result["char_count"], 0)
        self.assertEqual(result["sheet_count"], 0)
        self.assertEqual(result["row_count"], 0)
        self.assertEqual(result["cell_count"], 0)

    def test_unicode_preserved(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({
            "Թերթ": [[ARMENIAN, RUSSIAN, "Cargo transport"]],
        }))
        self.assertEqual(
            result["text"], f"[Sheet: Թերթ]\n{ARMENIAN}\t{RUSSIAN}\tCargo transport"
        )

    def test_numbers_bool_and_dates(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({
            "S": [[10, 2.5, True, False, datetime.datetime(2026, 9, 25, 18, 59, 34),
                   datetime.date(2026, 10, 2)]],
        }))
        self.assertEqual(
            result["text"],
            "[Sheet: S]\n10\t2.5\tTrue\tFalse\t2026-09-25T18:59:34\t2026-10-02T00:00:00",
        )

    def test_zero_is_not_treated_as_empty(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({"S": [[0, "x"]]}))
        self.assertEqual(result["text"], "[Sheet: S]\n0\tx")
        self.assertEqual(result["cell_count"], 2)

    def test_formula_is_not_evaluated(self):
        # openpyxl пишет формулу без закешированного значения: data_only даёт пустую ячейку.
        result = extractor.extract_xlsx(make_xlsx_bytes({"S": [["=1+1", "keep"]]}))
        self.assertEqual(result["text"], "[Sheet: S]\n\tkeep")

    def test_empty_rows_skipped_and_columns_aligned(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({
            "S": [["a", None, "c", None], [None, None], [], [None, "b"], ["", "  "]],
        }))
        # Хвостовые пустые ячейки отброшены, ведущие сохраняют выравнивание колонок.
        self.assertEqual(result["text"], "[Sheet: S]\na\t\tc\n\tb")
        self.assertEqual(result["row_count"], 2)
        self.assertEqual(result["cell_count"], 3)

    def test_strings_cleaned_minimally(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({"S": [["  a  b  ", "x\r\ny", "p\tq"]]}))
        self.assertEqual(result["text"], "[Sheet: S]\na  b\tx y\tp q")

    def test_char_count_matches_text_length(self):
        result = extractor.extract_xlsx(make_xlsx_bytes({"S": [[ARMENIAN, RUSSIAN]], "T": [[1]]}))
        self.assertEqual(result["char_count"], len(result["text"]))
        self.assertGreater(result["char_count"], 0)

    def test_bytes_file_like_and_path_sources(self):
        data = make_xlsx_bytes({"S": [["v"]]})
        expected = "[Sheet: S]\nv"
        self.assertEqual(extractor.extract_xlsx(data)["text"], expected)
        self.assertEqual(extractor.extract_xlsx(bytearray(data))["text"], expected)
        self.assertEqual(extractor.extract_xlsx(io.BytesIO(data))["text"], expected)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            path.write_bytes(data)
            self.assertEqual(extractor.extract_xlsx(path)["text"], expected)
            self.assertEqual(extractor.extract_xlsx(str(path))["text"], expected)

    def test_wrong_declared_dimension_does_not_lose_data(self):
        data = make_xlsx_bytes({"S": [["a", "b"], ["c", "d"]]})
        data = replace_zip_member(
            data, "xl/worksheets/sheet1.xml",
            lambda xml: xml.replace(b'<dimension ref="A1:B2"/>', b'<dimension ref="A1"/>'),
        )
        self.assertEqual(extractor.extract_xlsx(data)["text"], "[Sheet: S]\na\tb\nc\td")

    def test_opened_read_only_data_only_and_never_saved(self):
        data = make_xlsx_bytes({"S": [["v"]]})
        with mock.patch.object(
            extractor, "load_workbook", wraps=extractor.load_workbook
        ) as load, mock.patch("openpyxl.workbook.workbook.Workbook.save") as save:
            extractor.extract_xlsx(data)
        self.assertTrue(load.call_args.kwargs["read_only"])
        self.assertTrue(load.call_args.kwargs["data_only"])
        save.assert_not_called()

    def test_source_file_is_not_modified(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            path.write_bytes(make_xlsx_bytes({"S": [["v"]]}))
            before = (path.read_bytes(), path.stat().st_mtime_ns)
            extractor.extract_xlsx(path)
            self.assertEqual((path.read_bytes(), path.stat().st_mtime_ns), before)

    def test_workbook_is_closed_so_path_can_be_deleted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            path.write_bytes(make_xlsx_bytes({"S": [["v"]]}))
            extractor.extract_xlsx(path)
            path.unlink()  # на Windows упало бы, если бы файл остался открытым

    def test_broken_xlsx_raises(self):
        with self.assertRaises(zipfile.BadZipFile):
            extractor.extract_xlsx(b"this is not an xlsx file")

    def test_zip_without_workbook_raises(self):
        with self.assertRaises(Exception):
            extractor.extract_xlsx(make_zip_bytes([("hello.txt", b"hi")]))

    def test_default_limits(self):
        self.assertEqual(extractor.DEFAULT_XLSX_MAX_SHEETS, 100)
        self.assertEqual(extractor.DEFAULT_XLSX_MAX_ROWS, 100_000)
        self.assertEqual(extractor.DEFAULT_XLSX_MAX_CELLS, 1_000_000)

    def test_max_sheets(self):
        data = make_xlsx_bytes({f"S{i}": [["x"]] for i in range(3)})
        self.assertEqual(extractor.extract_xlsx(data, max_sheets=3)["sheet_count"], 3)
        with self.assertRaisesRegex(ValueError, "sheets"):
            extractor.extract_xlsx(data, max_sheets=2)

    def test_max_rows_counts_whole_workbook(self):
        data = make_xlsx_bytes({"A": [["1"], ["2"]], "B": [["3"], ["4"]]})
        self.assertEqual(extractor.extract_xlsx(data, max_rows=4)["row_count"], 4)
        with self.assertRaisesRegex(ValueError, "rows"):
            extractor.extract_xlsx(data, max_rows=3)

    def test_max_cells(self):
        data = make_xlsx_bytes({"S": [["a", "b", "c"], ["d", "e", "f"]]})
        self.assertEqual(extractor.extract_xlsx(data, max_cells=6)["cell_count"], 6)
        with self.assertRaisesRegex(ValueError, "cells"):
            extractor.extract_xlsx(data, max_cells=5)

    def test_max_chars(self):
        data = make_xlsx_bytes({"S": [["x" * 50]]})
        with self.assertRaisesRegex(ValueError, "characters"):
            extractor.extract_xlsx(data, max_chars=10)

    def test_limit_stops_reading_early(self):
        # Лимит срабатывает по мере чтения, а не после загрузки всех строк.
        data = make_xlsx_bytes({"S": [[i] for i in range(200)]})
        seen = []
        real = extractor._xlsx_row_cells

        def counting(row):
            seen.append(1)
            return real(row)

        with mock.patch.object(extractor, "_xlsx_row_cells", counting):
            with self.assertRaises(ValueError):
                extractor.extract_xlsx(data, max_rows=10)
        self.assertLessEqual(len(seen), 10)


class ExtractSupportedFromZipTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        self.zip_path = self.tmp / "tender.zip"

    def extract(self, members, **kwargs):
        make_zip(self.zip_path, members)
        return extractor.extract_supported_from_zip(self.zip_path, **kwargs)

    def test_docx_and_xlsx(self):
        result = self.extract([
            ("hraver.docx", make_docx_bytes([RUSSIAN])),
            ("spec.xlsx", make_xlsx_bytes({"Spec": [["Item", "Qty"], ["Bolt", 10]]})),
        ])
        self.assertEqual(result["file_type"], "zip")
        self.assertEqual(result["member_count"], 2)
        self.assertEqual(result["docx_count"], 1)
        self.assertEqual(result["xlsx_count"], 1)
        self.assertEqual(result["nested_zip_count"], 0)
        self.assertEqual(result["skipped_members"], [])
        self.assertEqual(result["failures"], [])

        docx, xlsx = result["documents"]
        self.assertEqual(docx["member_name"], "hraver.docx")
        self.assertEqual(docx["file_type"], "docx")
        self.assertEqual(docx["text"], RUSSIAN)
        self.assertEqual(docx["paragraph_count"], 1)
        self.assertEqual(xlsx["member_name"], "spec.xlsx")
        self.assertEqual(xlsx["file_type"], "xlsx")
        self.assertEqual(xlsx["text"], "[Sheet: Spec]\nItem\tQty\nBolt\t10")
        self.assertEqual(
            (xlsx["sheet_count"], xlsx["row_count"], xlsx["cell_count"]), (1, 2, 4)
        )
        self.assertEqual(xlsx["char_count"], len(xlsx["text"]))

    def test_30_xlsx_members(self):
        members = [
            (f"lot_{i}_i_texnikakan_bnutagir.xlsx", make_xlsx_bytes({"Lot": [[f"lot {i}", i]]}))
            for i in range(1, 31)
        ]
        result = self.extract(members)
        self.assertEqual(result["member_count"], 30)
        self.assertEqual(result["xlsx_count"], 30)
        self.assertEqual(len(result["documents"]), 30)
        self.assertEqual(result["skipped_members"], [])
        self.assertEqual(result["failures"], [])
        self.assertEqual(result["documents"][29]["text"], "[Sheet: Lot]\nlot 30\t30")

    def test_uppercase_xlsx_extension(self):
        result = self.extract([("SPEC.XLSX", make_xlsx_bytes({"S": [["v"]]}))])
        self.assertEqual(result["xlsx_count"], 1)

    def test_nested_zip_with_xlsx(self):
        inner = make_zip_bytes([("spec.xlsx", make_xlsx_bytes({"S": [["nested value"]]}))])
        result = self.extract([("lot_1_i_texnikakan_bnutagir.zip", inner)])
        self.assertEqual(result["member_count"], 1)
        self.assertEqual(result["nested_zip_count"], 1)
        self.assertEqual(result["xlsx_count"], 1)
        (document,) = result["documents"]
        self.assertEqual(document["member_name"], "lot_1_i_texnikakan_bnutagir.zip!/spec.xlsx")
        self.assertEqual(document["text"], "[Sheet: S]\nnested value")
        self.assertEqual(result["failures"], [])
        self.assertEqual(result["skipped_members"], [])

    def test_two_nested_zips_like_production(self):
        result = self.extract([
            (f"lot_{i}_i_texnikakan_bnutagir.zip",
             make_zip_bytes([(f"spec_{i}.xlsx", make_xlsx_bytes({"S": [[i]]}))]))
            for i in (1, 2)
        ])
        self.assertEqual(result["nested_zip_count"], 2)
        self.assertEqual(
            [d["member_name"] for d in result["documents"]],
            ["lot_1_i_texnikakan_bnutagir.zip!/spec_1.xlsx",
             "lot_2_i_texnikakan_bnutagir.zip!/spec_2.xlsx"],
        )

    def test_nested_zip_with_docx(self):
        inner = make_zip_bytes([("hraver.docx", make_docx_bytes([ARMENIAN]))])
        result = self.extract([("inner.zip", inner), ("top.docx", make_docx_bytes(["top"]))])
        self.assertEqual(result["docx_count"], 2)
        by_name = {d["member_name"]: d for d in result["documents"]}
        self.assertEqual(by_name["inner.zip!/hraver.docx"]["text"], ARMENIAN)
        self.assertEqual(by_name["top.docx"]["text"], "top")

    def test_nested_member_in_subfolder(self):
        inner = make_zip_bytes([("docs/spec.xlsx", make_xlsx_bytes({"S": [["v"]]}))])
        result = self.extract([("archive/inner.zip", inner)])
        self.assertEqual(result["documents"][0]["member_name"], "archive/inner.zip!/docs/spec.xlsx")

    def test_default_depth_allows_outer_and_one_nested_level(self):
        self.assertEqual(extractor.DEFAULT_MAX_NESTED_DEPTH, 2)
        deep = make_zip_bytes([("deep.xlsx", make_xlsx_bytes({"S": [["deep"]]}))])
        middle = make_zip_bytes([("ok.xlsx", make_xlsx_bytes({"S": [["ok"]]})), ("deep.zip", deep)])
        result = self.extract([("middle.zip", middle)])

        self.assertEqual([d["member_name"] for d in result["documents"]], ["middle.zip!/ok.xlsx"])
        # ZIP третьего уровня не читается, но попадает в skipped_members.
        self.assertEqual(result["skipped_members"], ["middle.zip!/deep.zip"])
        self.assertEqual(result["failures"], [])
        self.assertEqual(result["nested_zip_count"], 1)

    def test_max_nested_depth_is_configurable(self):
        deep = make_zip_bytes([("deep.xlsx", make_xlsx_bytes({"S": [["deep"]]}))])
        middle = make_zip_bytes([("deep.zip", deep)])
        members = [("middle.zip", middle)]

        result = self.extract(members, max_nested_depth=3)
        self.assertEqual([d["member_name"] for d in result["documents"]],
                         ["middle.zip!/deep.zip!/deep.xlsx"])
        self.assertEqual(result["skipped_members"], [])

        result = self.extract(members, max_nested_depth=1)
        self.assertEqual(result["documents"], [])
        self.assertEqual(result["skipped_members"], ["middle.zip"])

    def test_deeply_nested_chain_is_not_followed(self):
        payload = make_zip_bytes([("x.xlsx", make_xlsx_bytes({"S": [["x"]]}))])
        for _ in range(20):
            payload = make_zip_bytes([("z.zip", payload)])
        result = self.extract([("z.zip", payload)])
        self.assertEqual(result["documents"], [])
        self.assertEqual(len(result["skipped_members"]), 1)

    def test_broken_nested_zip_is_isolated(self):
        good = make_zip_bytes([("spec.xlsx", make_xlsx_bytes({"S": [["good"]]}))])
        result = self.extract([
            ("broken.zip", b"this is not a zip"),
            ("good.zip", good),
            ("top.docx", make_docx_bytes(["top"])),
        ])
        self.assertEqual(
            [d["member_name"] for d in result["documents"]],
            ["good.zip!/spec.xlsx", "top.docx"],
        )
        self.assertEqual(result["nested_zip_count"], 1)
        (failure,) = result["failures"]
        self.assertEqual(failure["member_name"], "broken.zip")
        self.assertEqual(failure["error_type"], "BadZipFile")
        self.assertTrue(failure["error_message"])

    def test_broken_xlsx_and_docx_are_isolated_including_nested(self):
        inner = make_zip_bytes([
            ("bad.xlsx", b"not xlsx"),
            ("good.xlsx", make_xlsx_bytes({"S": [["good"]]})),
        ])
        result = self.extract([
            ("bad.docx", b"not docx"),
            ("inner.zip", inner),
        ])
        self.assertEqual([d["member_name"] for d in result["documents"]], ["inner.zip!/good.xlsx"])
        self.assertEqual(
            [f["member_name"] for f in result["failures"]], ["bad.docx", "inner.zip!/bad.xlsx"]
        )

    def test_unsupported_members_skipped_at_every_level(self):
        inner = make_zip_bytes([
            ("scan.pdf", b"%PDF-1.4"),
            ("photo.jpg", b"\xff\xd8"),
            ("spec.xlsx", make_xlsx_bytes({"S": [["v"]]})),
        ])
        result = self.extract([
            ("archive.rar", b"Rar!\x1a\x07\x00"),
            ("data.xml", b"<a/>"),
            ("old.doc", b"\xd0\xcf"),
            ("old.xls", b"\xd0\xcf"),
            ("inner.zip", inner),
        ])
        self.assertEqual(result["xlsx_count"], 1)
        self.assertEqual(result["failures"], [])
        self.assertEqual(result["skipped_members"], [
            "archive.rar", "data.xml", "old.doc", "old.xls",
            "inner.zip!/scan.pdf", "inner.zip!/photo.jpg",
        ])

    def test_rar_stays_skipped_not_failure(self):
        result = self.extract([("docs.rar", b"Rar!\x1a\x07\x00 payload")])
        self.assertEqual(result["skipped_members"], ["docs.rar"])
        self.assertEqual(result["failures"], [])
        self.assertEqual(result["documents"], [])

    def test_member_count_budget_shared_with_nested_zip(self):
        def inner():
            return make_zip_bytes([(f"{i}.txt", b"x") for i in range(3)])

        members = [("a.zip", inner()), ("b.zip", inner())]
        # outer 2 + 3 + 3 = 8 member: каждый архив по отдельности укладывается в 6.
        with self.assertRaisesRegex(ValueError, "too many members"):
            self.extract(members, max_members=6)
        self.assertEqual(self.extract(members, max_members=8)["member_count"], 2)

    def test_total_size_budget_cannot_be_bypassed_through_nested_zip(self):
        def inner():
            return make_zip_bytes(
                [("zeros.bin", b"\0" * 100_000)], compression=zipfile.ZIP_DEFLATED
            )

        members = [("a.zip", inner()), ("b.zip", inner())]
        # Каждый вложенный ZIP (100 000 байт) по отдельности меньше лимита 150 000,
        # но вместе они его превышают.
        with self.assertRaisesRegex(ValueError, "total uncompressed size"):
            self.extract(members, max_total_uncompressed_bytes=150_000)
        self.extract(members[:1], max_total_uncompressed_bytes=150_000)

    def test_member_size_limit_applies_to_nested_members(self):
        inner = make_zip_bytes([("big.bin", b"x" * 500)], compression=zipfile.ZIP_DEFLATED)
        with self.assertRaisesRegex(ValueError, "too large"):
            self.extract([("inner.zip", inner)], max_member_bytes=400)

    def test_nested_limit_violation_is_not_swallowed_as_failure(self):
        inner = make_zip_bytes([(f"{i}.txt", b"x") for i in range(10)])
        good = make_docx_bytes(["good"])
        with self.assertRaises(extractor.ZipLimitError):
            self.extract([("good.docx", good), ("inner.zip", inner)], max_members=5)

    def test_limit_error_is_value_error(self):
        self.assertTrue(issubclass(extractor.ZipLimitError, ValueError))

    def test_nested_zip_writes_nothing_to_filesystem(self):
        work = self.tmp / "work"
        work.mkdir()
        zip_path = work / "tender.zip"
        inner = make_zip_bytes([("../../evil.xlsx", make_xlsx_bytes({"S": [["evil"]]}))])
        make_zip(zip_path, [("../inner.zip", inner)])
        before_tree = sorted(str(p) for p in self.tmp.rglob("*"))
        cwd_before = sorted(os.listdir("."))

        result = extractor.extract_supported_from_zip(zip_path)

        self.assertEqual(result["documents"][0]["member_name"], "../inner.zip!/../../evil.xlsx")
        self.assertEqual(sorted(str(p) for p in self.tmp.rglob("*")), before_tree)
        self.assertEqual(sorted(os.listdir(".")), cwd_before)

    def test_bytes_and_file_like_sources(self):
        data = make_zip_bytes([("spec.xlsx", make_xlsx_bytes({"S": [["v"]]}))])
        for source in (data, io.BytesIO(data)):
            result = extractor.extract_supported_from_zip(source)
            self.assertEqual(result["documents"][0]["text"], "[Sheet: S]\nv")

    def test_extract_docx_from_zip_is_backward_compatible_wrapper(self):
        inner = make_zip_bytes([("spec.xlsx", make_xlsx_bytes({"S": [["v"]]}))])
        make_zip(self.zip_path, [("a.docx", make_docx_bytes(["a"])), ("inner.zip", inner)])
        old = extractor.extract_docx_from_zip(self.zip_path)
        self.assertEqual(old, extractor.extract_supported_from_zip(self.zip_path))
        self.assertEqual(old["docx_count"], 1)
        self.assertEqual(old["xlsx_count"], 1)


if __name__ == "__main__":
    unittest.main()
