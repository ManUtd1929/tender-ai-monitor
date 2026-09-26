"""
Production-модуль извлечения текста из тендерных документов.

Поддерживается только:
    - .docx (путь, file-like объект или bytes);
    - .xlsx (путь, file-like объект или bytes);
    - .zip с .docx, .xlsx и вложенными .zip (в памяти, с ограничением глубины).

НЕ поддерживается: .doc, .pdf, .xls, .rar, OCR, AI. В monitor.py и enrichment
pipeline модуль НЕ подключён; ничего не читается при импорте и без явного
вызова функций. HTTP не используется.

extract_docx() возвращает dict:

    file_type, text, char_count, paragraph_count, table_count

Текст собирается в порядке документа: непустые абзацы и строки таблиц идут
построчно, разделитель — "\n". Строка таблицы: ячейки через TAB. Строки, в
которых все ячейки пустые, пропускаются. Таблицы не превращаются в структуру:
нужен читаемый текст для будущего AI. Вложенные таблицы, колонтитулы и
содержимое content controls (w:sdt) не извлекаются.
paragraph_count — число непустых абзацев верхнего уровня (вне таблиц),
table_count — число таблиц верхнего уровня.

extract_xlsx() возвращает dict:

    file_type, text, char_count, sheet_count, row_count, cell_count

Книга открывается openpyxl только на чтение (read_only=True, data_only=True) и
никогда не сохраняется. Формулы не вычисляются: берётся значение, которое Excel
закешировал в файле (если кеша нет — пустая ячейка). Формат текста:

    [Sheet: <название листа>]
    значение<TAB>значение<TAB>значение
    ...

Листы разделены пустой строкой. Полностью пустые листы и строки пропускаются,
хвостовые пустые ячейки строки отбрасываются. sheet_count / row_count /
cell_count считают только непустые листы, строки и ячейки. Переводы строки и TAB
внутри ячейки заменяются пробелом, чтобы не ломать построчную структуру.
Лимиты (листы, строки, ячейки, символы) проверяются по фактически прочитанным
данным; превышение -> ValueError. Повреждённый XLSX приводит к исключению.

Повреждённый DOCX в extract_docx() приводит к исключению (python-docx /
zipfile), оно не скрывается.

extract_supported_from_zip() читает ZIP в памяти, на диск ничего не распаковывает
и не использует имена member как пути файловой системы: member_name — только
метаданные. Вложенный ZIP имеет member_name "<внешний member>!/<внутренний member>".
max_nested_depth — максимальное число уровней ZIP, считая внешний: при значении
по умолчанию (2) читаются outer.zip -> inner.zip -> docx/xlsx, а ZIP третьего уровня
попадает в skipped_members. Лимиты (число member, размер member, суммарный размер) —
единый бюджет на всю рекурсию: они проверяются по ZIP-заголовкам ДО чтения и ещё раз
по фактически прочитанным bytes; вложенный ZIP учитывается и как member внешнего, и
через свои member. Превышение -> ValueError для всего результата. Повреждённый
DOCX/XLSX/вложенный ZIP не останавливает обработку: он попадает в failures, остальные
member обрабатываются. extract_docx_from_zip() — прежнее имя (обёртка).
"""

import datetime
import io
import logging
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

DOCX_EXTENSION = ".docx"
XLSX_EXTENSION = ".xlsx"
ZIP_EXTENSION = ".zip"

DEFAULT_MAX_MEMBERS = 200
DEFAULT_MAX_MEMBER_BYTES = 20 * 1024 * 1024
DEFAULT_MAX_TOTAL_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
DEFAULT_MAX_NESTED_DEPTH = 2

DEFAULT_XLSX_MAX_SHEETS = 100
DEFAULT_XLSX_MAX_ROWS = 100_000
DEFAULT_XLSX_MAX_CELLS = 1_000_000
DEFAULT_XLSX_MAX_CHARS = 10_000_000

NESTED_MEMBER_DELIMITER = "!/"


def clean_text_piece(value):
    """strip() и нормализация CRLF/CR -> LF. Больше ничего не меняется."""
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def _open_arg(source):
    """Приводит source к тому, что принимают Document() и ZipFile()."""
    if isinstance(source, (bytes, bytearray)):
        return io.BytesIO(source)
    if isinstance(source, Path):
        return str(source)
    return source


def _cell_text(cell):
    parts = (clean_text_piece(p.text) for p in cell.paragraphs)
    return " ".join(p for p in parts if p)


def _table_lines(table):
    lines = []
    for row in table.rows:
        # Ячейки берём напрямую из XML строки: row.cells повторяет объединённую
        # ячейку для каждой колонки/строки, что дублировало бы текст.
        cells = [_cell_text(_Cell(tc, table)) for tc in row._tr.tc_lst]
        if any(cells):
            lines.append("\t".join(cells))
    return lines


def extract_docx(source):
    """Извлекает текст абзацев и таблиц из .docx (path / file-like / bytes)."""
    document = Document(_open_arg(source))

    lines = []
    paragraph_count = 0
    table_count = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = clean_text_piece(Paragraph(child, document).text)
            if text:
                lines.append(text)
                paragraph_count += 1
        elif child.tag == qn("w:tbl"):
            table_count += 1
            lines.extend(_table_lines(Table(child, document)))

    text = "\n".join(lines)
    return {
        "file_type": "docx",
        "text": text,
        "char_count": len(text),
        "paragraph_count": paragraph_count,
        "table_count": table_count,
    }




# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------

def _xlsx_cell_text(value):
    """Ячейка -> строка. Формулы не вычисляются: приходит закешированное значение."""
    if value is None:
        return ""
    if isinstance(value, str):
        text = clean_text_piece(value)
    elif isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        text = value.isoformat()
    else:
        text = str(value)
    # Перевод строки и TAB внутри ячейки сломали бы построчную структуру.
    return text.replace("\n", " ").replace("\t", " ")


def _xlsx_row_cells(row):
    """Значения строки; хвостовые пустые ячейки отбрасываются."""
    cells = [_xlsx_cell_text(value) for value in row]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _xlsx_limit_error(what, count, limit):
    return ValueError(f"XLSX has too many {what}: {count} > {limit}")


def extract_xlsx(
    source,
    max_sheets=DEFAULT_XLSX_MAX_SHEETS,
    max_rows=DEFAULT_XLSX_MAX_ROWS,
    max_cells=DEFAULT_XLSX_MAX_CELLS,
    max_chars=DEFAULT_XLSX_MAX_CHARS,
):
    """
    Извлекает текст всех непустых листов из .xlsx (path / file-like / bytes).

    max_rows и max_cells считаются по фактически прочитанным строкам и ячейкам
    (включая пустые) всей книги, max_chars — по накопленному тексту.
    """
    workbook = load_workbook(
        _open_arg(source), read_only=True, data_only=True, keep_links=False
    )
    try:
        worksheets = workbook.worksheets
        if len(worksheets) > max_sheets:
            raise _xlsx_limit_error("sheets", len(worksheets), max_sheets)

        sections = []
        sheet_count = row_count = cell_count = 0
        rows_read = cells_read = chars = 0

        for worksheet in worksheets:
            # Заявленный <dimension> в файле часто неверен; без сброса read-only
            # режим может отдать не все ячейки или пустые строки на всю ширину листа.
            worksheet.reset_dimensions()

            lines = []
            for row in worksheet.iter_rows(values_only=True):
                rows_read += 1
                if rows_read > max_rows:
                    raise _xlsx_limit_error("rows", rows_read, max_rows)
                cells_read += len(row)
                if cells_read > max_cells:
                    raise _xlsx_limit_error("cells", cells_read, max_cells)

                cells = _xlsx_row_cells(row)
                if not cells:
                    continue
                line = "\t".join(cells)
                chars += len(line) + 1
                if chars > max_chars:
                    raise _xlsx_limit_error("characters", chars, max_chars)
                lines.append(line)
                row_count += 1
                cell_count += sum(1 for cell in cells if cell)

            if lines:
                sheet_count += 1
                sections.append("\n".join([f"[Sheet: {worksheet.title}]", *lines]))
    finally:
        workbook.close()

    text = "\n\n".join(sections)
    return {
        "file_type": "xlsx",
        "text": text,
        "char_count": len(text),
        "sheet_count": sheet_count,
        "row_count": row_count,
        "cell_count": cell_count,
    }


# --------------------------------------------------------------------------
# ZIP
# --------------------------------------------------------------------------

class ZipLimitError(ValueError):
    """Превышен лимит ZIP. Не изолируется по member: прерывает весь результат."""


class _ZipBudget:
    """Единые лимиты на внешний ZIP и все вложенные: вложенность их не обходит."""

    def __init__(self, max_members, max_member_bytes, max_total_bytes):
        self.max_members = max_members
        self.max_member_bytes = max_member_bytes
        self.max_total_bytes = max_total_bytes
        self.declared_members = 0
        self.declared_bytes = 0
        self.read_bytes = 0

    def check_archive(self, members):
        """Проверка по заголовкам ZIP до чтения содержимого."""
        self.declared_members += len(members)
        if self.declared_members > self.max_members:
            raise ZipLimitError(
                f"ZIP has too many members: {self.declared_members} > {self.max_members}"
            )
        for info in members:
            if info.file_size > self.max_member_bytes:
                raise ZipLimitError(
                    f"ZIP member {info.filename!r} is too large: "
                    f"{info.file_size} > {self.max_member_bytes} bytes"
                )
            self.declared_bytes += info.file_size
        if self.declared_bytes > self.max_total_bytes:
            raise ZipLimitError(
                "ZIP total uncompressed size is too large: "
                f"{self.declared_bytes} > {self.max_total_bytes} bytes"
            )

    def read_member(self, zf, info):
        """Читает member; лимиты проверяются ещё раз по реально прочитанным bytes."""
        with zf.open(info) as member_file:
            # +1 байт: заголовок file_size мог занижать реальный размер.
            data = member_file.read(self.max_member_bytes + 1)
        if len(data) > self.max_member_bytes:
            raise ZipLimitError(
                f"ZIP member {info.filename!r} is too large: > {self.max_member_bytes} bytes"
            )
        self.read_bytes += len(data)
        if self.read_bytes > self.max_total_bytes:
            raise ZipLimitError(
                f"ZIP total uncompressed size is too large: > {self.max_total_bytes} bytes"
            )
        return data


def _failure(member_name, exc):
    return {
        "member_name": member_name,
        "error_type": type(exc).__name__,
        "error_message": str(exc),
    }


def _member_extension(name):
    lowered = name.lower()
    for extension in (DOCX_EXTENSION, XLSX_EXTENSION, ZIP_EXTENSION):
        if lowered.endswith(extension):
            return extension
    return None


def _walk_zip(source, prefix, level, budget, max_nested_depth, found):
    """
    Читает один ZIP (level 1 — внешний) и дополняет found. Ошибки открытия и
    превышение лимитов пробрасываются; повреждённый member обрабатывается здесь.
    """
    with zipfile.ZipFile(_open_arg(source)) as zf:
        members = [info for info in zf.infolist() if not info.is_dir()]
        budget.check_archive(members)
        if level == 1:
            found["member_count"] = len(members)

        for info in members:
            name = prefix + info.filename
            extension = _member_extension(info.filename)

            if extension is None:
                logger.info("ZIP member skipped (unsupported type): %s", name)
                found["skipped_members"].append(name)
                continue
            if extension == ZIP_EXTENSION and level >= max_nested_depth:
                logger.warning(
                    "Nested ZIP skipped (max_nested_depth=%d): %s", max_nested_depth, name
                )
                found["skipped_members"].append(name)
                continue

            try:
                data = budget.read_member(zf, info)
            except ZipLimitError:
                raise
            except Exception as exc:
                logger.warning(
                    "ZIP member read failed: %s (%s: %s)", name, type(exc).__name__, exc
                )
                found["failures"].append(_failure(name, exc))
                continue

            if extension == ZIP_EXTENSION:
                try:
                    _walk_zip(
                        data, name + NESTED_MEMBER_DELIMITER, level + 1,
                        budget, max_nested_depth, found,
                    )
                except ZipLimitError:
                    raise
                except Exception as exc:
                    logger.warning(
                        "Nested ZIP failed: %s (%s: %s)", name, type(exc).__name__, exc
                    )
                    found["failures"].append(_failure(name, exc))
                else:
                    found["nested_zip_count"] += 1
                    logger.info("Nested ZIP processed: %s", name)
                continue

            kind = extension[1:].upper()
            extract = extract_docx if extension == DOCX_EXTENSION else extract_xlsx
            try:
                result = extract(data)
            except Exception as exc:
                logger.warning(
                    "%s extraction failed: %s (%s: %s)", kind, name, type(exc).__name__, exc
                )
                found["failures"].append(_failure(name, exc))
                continue

            logger.info("%s extracted: %s (%d chars)", kind, name, result["char_count"])
            found["documents"].append({"member_name": name, **result})


def extract_supported_from_zip(
    zip_path,
    max_members=DEFAULT_MAX_MEMBERS,
    max_member_bytes=DEFAULT_MAX_MEMBER_BYTES,
    max_total_uncompressed_bytes=DEFAULT_MAX_TOTAL_UNCOMPRESSED_BYTES,
    max_nested_depth=DEFAULT_MAX_NESTED_DEPTH,
):
    """
    Извлекает текст из всех .docx и .xlsx внутри ZIP (в том числе во вложенных
    ZIP) без распаковки на диск.

    member_count — число member внешнего ZIP без директорий; nested_zip_count —
    число успешно прочитанных вложенных ZIP. documents — успешные извлечения со
    всех уровней: {member_name, file_type, ...поля extract_docx / extract_xlsx}.
    skipped_members — имена неподдерживаемых member (и ZIP глубже max_nested_depth).
    failures — повреждённые DOCX/XLSX/вложенные ZIP и member, которые не удалось
    прочитать: {member_name, error_type, error_message}.
    """
    found = {
        "member_count": 0,
        "nested_zip_count": 0,
        "documents": [],
        "skipped_members": [],
        "failures": [],
    }
    budget = _ZipBudget(max_members, max_member_bytes, max_total_uncompressed_bytes)
    _walk_zip(zip_path, "", 1, budget, max_nested_depth, found)

    documents = found["documents"]
    return {
        "file_type": "zip",
        "member_count": found["member_count"],
        "nested_zip_count": found["nested_zip_count"],
        "docx_count": sum(1 for item in documents if item["file_type"] == "docx"),
        "xlsx_count": sum(1 for item in documents if item["file_type"] == "xlsx"),
        "documents": documents,
        "skipped_members": found["skipped_members"],
        "failures": found["failures"],
    }


def extract_docx_from_zip(
    zip_path,
    max_members=DEFAULT_MAX_MEMBERS,
    max_member_bytes=DEFAULT_MAX_MEMBER_BYTES,
    max_total_uncompressed_bytes=DEFAULT_MAX_TOTAL_UNCOMPRESSED_BYTES,
    max_nested_depth=DEFAULT_MAX_NESTED_DEPTH,
):
    """Прежнее имя extract_supported_from_zip(): результат содержит и .xlsx, и вложенные ZIP."""
    return extract_supported_from_zip(
        zip_path,
        max_members=max_members,
        max_member_bytes=max_member_bytes,
        max_total_uncompressed_bytes=max_total_uncompressed_bytes,
        max_nested_depth=max_nested_depth,
    )


# --------------------------------------------------------------------------

def main():
    print("Document extractor module. Use explicit extraction functions.")


if __name__ == "__main__":
    main()
